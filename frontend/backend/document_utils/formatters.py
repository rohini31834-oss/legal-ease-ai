from __future__ import annotations

import io
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from fpdf import FPDF
from fpdf.enums import XPos, YPos

def clean_filename(
    name: str,
    extension: str
) -> str:

    stem = Path(
        name or "legalease_document"
    ).stem

    stem = re.sub(
        r"[^A-Za-z0-9._-]+",
        "_",
        stem
    )

    stem = stem.strip("._")

    if not stem:
        stem = "legalease_document"

    return f"{stem}.{extension}"


def document_to_text(
    markdown: str
) -> str:

    text = markdown.replace(
        "\r\n",
        "\n"
    )

    text = re.sub(
        r"!\[[^\]]*\]\([^)]+\)",
        "",
        text
    )

    text = re.sub(
        r"\[([^\]]+)\]\([^)]+\)",
        r"\1",
        text
    )

    text = re.sub(
        r"^#{1,6}\s*",
        "",
        text,
        flags=re.MULTILINE
    )

    text = re.sub(
        r"\*\*(.*?)\*\*",
        r"\1",
        text
    )

    text = re.sub(
        r"__(.*?)__",
        r"\1",
        text
    )

    text = re.sub(
        r"`(.*?)`",
        r"\1",
        text
    )

    text = re.sub(
        r"^\s*[-*]\s+",
        "• ",
        text,
        flags=re.MULTILINE
    )

    return text.strip()


def build_docx(
    title: str,
    markdown: str
) -> io.BytesIO:

    doc = Document()

    section = doc.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]

    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)

    heading = doc.add_heading(
        title or "Legal Document",
        level=0
    )

    heading.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    for raw in markdown.splitlines():

        line = raw.strip()

        if not line:

            doc.add_paragraph("")

            continue

        if line.startswith("### "):

            doc.add_heading(
                line[4:],
                level=2
            )

        elif line.startswith("## "):

            doc.add_heading(
                line[3:],
                level=1
            )

        elif line.startswith("# "):

            doc.add_heading(
                line[2:],
                level=1
            )

        elif re.match(
            r"^\d+\.\s+",
            line
        ):

            doc.add_paragraph(
                line,
                style="List Number"
            )

        elif re.match(
            r"^[-*]\s+",
            line
        ):

            doc.add_paragraph(
                re.sub(
                    r"^[-*]\s+",
                    "",
                    line
                ),
                style="List Bullet"
            )

        else:

            paragraph = doc.add_paragraph()

            _add_markdown_runs(
                paragraph,
                line
            )

    footer = (
        doc.sections[0]
        .footer
        .paragraphs[0]
    )

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer.add_run(
        "Generated with LegalEaseAI — Review before use."
    )

    buffer = io.BytesIO()

    doc.save(buffer)

    buffer.seek(0)

    return buffer


def _add_markdown_runs(
    paragraph,
    text: str
) -> None:

    pattern = re.compile(
        r"(\*\*.*?\*\*|__.*?__)"
    )

    parts = pattern.split(text)

    for part in parts:

        if not part:
            continue

        if (
            (
                part.startswith("**")
                and part.endswith("**")
            )
            or
            (
                part.startswith("__")
                and part.endswith("__")
            )
        ):

            run = paragraph.add_run(
                part[2:-2]
            )

            run.bold = True

        else:

            paragraph.add_run(part)


class LegalPDF(FPDF):

    def header(self):

        self.set_font(
            "Helvetica",
            "B",
            9
        )

        self.cell(
            0,
            7,
            "LegalEaseAI",
            align="L"
        )

        self.ln(5)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "", 8)

        footer_text = (
            f"Page {self.page_no()} | "
            "AI-generated draft - review before use."
        )

        self.cell(
            0,
            8,
            _pdf_safe(footer_text),
            align="C"
        )


def build_pdf(
    title: str,
    markdown: str
) -> io.BytesIO:

    pdf = LegalPDF()

    pdf.set_auto_page_break(
        auto=True,
        margin=18
    )

    pdf.set_margins(
        18,
        18,
        18
    )

    pdf.add_page()

    pdf.set_font(
        "Helvetica",
        "B",
        16
    )

    pdf.multi_cell(
    0,
    9,
    _pdf_safe(
        title or "Legal Document"
    ),
    align="C",
    new_x="LMARGIN",
    new_y="NEXT"
)

    pdf.ln(4)

    for raw in markdown.splitlines():
        line = raw.strip()

        if not line:
            pdf.ln(3)
            continue

        if line.startswith("#"):
            heading = re.sub(r"^#+\s*", "", line)
            pdf.set_font("Helvetica", "B", 12)
            pdf.multi_cell(
    0,
    7,
    _pdf_safe(
        _strip_markdown(heading)
    ),
    new_x="LMARGIN",
    new_y="NEXT"
)
            pdf.ln(1)
        else:
            pdf.set_font("Helvetica", "", 10)
            pdf.multi_cell(
    0,
    5.5,
    _pdf_safe(
        _strip_markdown(line)
    ),
    new_x="LMARGIN",
    new_y="NEXT"
)
    pdf_data = pdf.output()

    output = io.BytesIO(bytes(pdf_data))
    output.seek(0)

    return output


def _strip_markdown(
    text: str
) -> str:

    text = re.sub(
        r"\*\*(.*?)\*\*",
        r"\1",
        text
    )

    text = re.sub(
        r"__(.*?)__",
        r"\1",
        text
    )

    text = re.sub(
        r"`(.*?)`",
        r"\1",
        text
    )

    return text


def _pdf_safe(
    text: str
) -> str:

    replacements = {
        "—": "-",
        "–": "-",
        "“": '"',
        "”": '"',
        "‘": "'",
        "’": "'",
        "₹": "Rs.",
        "•": "-",
        "\u00a0": " ",
    }

    for old, new in replacements.items():

        text = text.replace(
            old,
            new
        )

    return (
        text
        .encode(
            "latin-1",
            "replace"
        )
        .decode("latin-1")
    )